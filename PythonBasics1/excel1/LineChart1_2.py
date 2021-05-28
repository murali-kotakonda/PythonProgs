import random
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl import load_workbook

workbook = workbook = load_workbook('line_chart.xlsx')
sheet = workbook.active

"""
1.create the chart obj
2.to the chart obj add the data to be displyed
3.add the chart obj to the sheet
"""
#create the chart obj
chart = LineChart()

#create the refernce to be added to the chart
data = Reference(worksheet=sheet,
                 min_row=2,
                 max_row=4,
                 min_col=1,
                 max_col=13)

#add the data to the cart
chart.add_data(data, from_rows=True, titles_from_data=True)

#add the chart to the sheet
sheet.add_chart(chart, "C6")

workbook.save("line_chart.xlsx")